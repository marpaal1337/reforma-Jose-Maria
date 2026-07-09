#!/usr/bin/env python3
"""Lee Presupuesto/Presupuesto.xlsx y lo vuelca a data/excel.json."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Presupuesto" / "Presupuesto.xlsx"
OUT = ROOT / "data" / "excel.json"


def main() -> None:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    sheets = {}
    for s in wb.sheetnames:
        ws = wb[s]
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
        sheets[s] = {
            "dim": [ws.max_row, ws.max_column],
            "filas": rows,
        }
    OUT.write_text(
        json.dumps({"archivo": str(SRC.relative_to(ROOT)), "hojas": sheets},
                   ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
